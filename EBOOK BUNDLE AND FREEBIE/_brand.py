"""
Shared brand helpers for Six & Thriving freebies (PowerPoint editions).
Imported by all premium .pptx generators in this folder.
"""

from pptx.util import Pt, Cm
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ── BRAND PALETTE ────────────────────────────────────────────────────
NAVY = RGBColor(0x0D, 0x1B, 0x3E)
NAVY_MID = RGBColor(0x16, 0x23, 0x47)
NAVY_LITE = RGBColor(0x1E, 0x2F, 0x55)
GOLD = RGBColor(0xC9, 0xA8, 0x4C)
GOLD_LITE = RGBColor(0xE8, 0xC9, 0x7A)
CREAM = RGBColor(0xFA, 0xF7, 0xF2)
CREAM_MID = RGBColor(0xF0, 0xEB, 0xE1)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TEAL = RGBColor(0x2D, 0x8A, 0x8A)
SLATE = RGBColor(0x4A, 0x55, 0x68)
SLATE_LITE = RGBColor(0x71, 0x80, 0x96)
RUST = RGBColor(0xC0, 0x39, 0x2B)
GREEN_DK = RGBColor(0x27, 0x67, 0x49)
AMBER = RGBColor(0xD6, 0x9E, 0x2E)
PURPLE = RGBColor(0x55, 0x3C, 0x9A)
AMBER_LT = RGBColor(0xFE, 0xFC, 0xBF)
BLUE_LT = RGBColor(0xEB, 0xF8, 0xFF)
RED_LT = RGBColor(0xFE, 0xD7, 0xD7)
GREEN_LT = RGBColor(0xF0, 0xFD, 0xF4)

FONT_HEADING = "Georgia"
FONT_BODY = "Arial"

# A4 portrait — matches the main ebook
SW = Cm(21.0)
SH = Cm(29.7)


# ── HELPERS ───────────────────────────────────────────────────────────

def bg(slide, color):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = color


def shp(slide, l, t, w, h, color, st=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(st, l, t, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    return s


def txt(slide, l, t, w, h, text, sz=Pt(10), clr=SLATE, b=False, i=False,
        al=PP_ALIGN.LEFT, fn=FONT_BODY, anc=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anc
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = sz
    p.font.color.rgb = clr
    p.font.bold = b
    p.font.italic = i
    p.font.name = fn
    p.alignment = al
    return tb


def mtxt(slide, l, t, w, h, paras):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for idx, pd in enumerate(paras):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = pd.get("t", "")
        p.font.size = pd.get("sz", Pt(10))
        p.font.color.rgb = pd.get("c", SLATE)
        p.font.bold = pd.get("b", False)
        p.font.italic = pd.get("i", False)
        p.font.name = pd.get("fn", FONT_BODY)
        p.alignment = pd.get("al", PP_ALIGN.LEFT)
        p.space_after = pd.get("sa", Pt(6))
    return tb


def gbar(slide, l, t, w):
    return shp(slide, l, t, w, Cm(0.12), GOLD, MSO_SHAPE.RECTANGLE)


def header(slide, label_right):
    """Standard navy header band."""
    shp(slide, Cm(0), Cm(0), SW, Cm(1.3), NAVY, MSO_SHAPE.RECTANGLE)
    shp(slide, Cm(0), Cm(1.3), SW, Cm(0.12), GOLD, MSO_SHAPE.RECTANGLE)
    txt(slide, Cm(2), Cm(0.3), Cm(8), Cm(0.7),
        "Sleep, Baby. Please.", Pt(7.5), CREAM, fn=FONT_HEADING)
    txt(slide, Cm(11), Cm(0.3), Cm(8), Cm(0.7),
        label_right, Pt(7.5), GOLD_LITE, al=PP_ALIGN.RIGHT, fn=FONT_BODY)


def footer(slide, pg):
    """Cream footer with navy page-number circle."""
    shp(slide, Cm(0), Cm(28.1), SW, Cm(1.6), CREAM_MID, MSO_SHAPE.RECTANGLE)
    shp(slide, Cm(0), Cm(28.0), SW, Cm(0.1), GOLD, MSO_SHAPE.RECTANGLE)
    shp(slide, Cm(9.6), Cm(28.3), Cm(1.3), Cm(1.3), NAVY, MSO_SHAPE.OVAL)
    txt(slide, Cm(9.6), Cm(28.4), Cm(1.3), Cm(1.1), str(pg),
        Pt(9), WHITE, b=True, al=PP_ALIGN.CENTER, anc=MSO_ANCHOR.MIDDLE)
    txt(slide, Cm(2), Cm(28.4), Cm(7), Cm(0.7),
        "Six & Thriving  •  www.sixandthriving.com", Pt(7), SLATE_LITE)
    txt(slide, Cm(12.5), Cm(28.4), Cm(6.5), Cm(0.7),
        "© 2026 Six & Thriving", Pt(7), SLATE_LITE, al=PP_ALIGN.RIGHT)


def dark_footer(slide, pg):
    shp(slide, Cm(0), Cm(28.1), SW, Cm(1.6), RGBColor(0x08, 0x0F, 0x24), MSO_SHAPE.RECTANGLE)
    shp(slide, Cm(0), Cm(28.0), SW, Cm(0.1), GOLD, MSO_SHAPE.RECTANGLE)
    shp(slide, Cm(9.6), Cm(28.3), Cm(1.3), Cm(1.3), GOLD, MSO_SHAPE.OVAL)
    txt(slide, Cm(9.6), Cm(28.4), Cm(1.3), Cm(1.1), str(pg),
        Pt(9), NAVY, b=True, al=PP_ALIGN.CENTER, anc=MSO_ANCHOR.MIDDLE)
    txt(slide, Cm(2), Cm(28.4), Cm(7), Cm(0.7),
        "www.sixandthriving.com", Pt(7), GOLD_LITE)
    txt(slide, Cm(12.5), Cm(28.4), Cm(6.5), Cm(0.7),
        "© 2026 Six & Thriving", Pt(7), SLATE_LITE, al=PP_ALIGN.RIGHT)


def callout(slide, top, label, body, accent=TEAL, bg_clr=BLUE_LT,
            text_color=None, h=None):
    """Reusable callout box with accent strip."""
    if h is None:
        h = Cm(2.5)
    shp(slide, Cm(2), top, Cm(17), h, bg_clr, MSO_SHAPE.ROUNDED_RECTANGLE)
    shp(slide, Cm(2), top, Cm(0.2), h, accent, MSO_SHAPE.RECTANGLE)
    txt(slide, Cm(2.6), top + Cm(0.2), Cm(16), Cm(0.4),
        label, Pt(7.5), accent, b=True)
    txt(slide, Cm(2.6), top + Cm(0.7), Cm(16), h - Cm(0.9),
        body, Pt(9.5), text_color or NAVY)


def cta_box(slide, top, title, subtitle, button_text, url):
    """Standard 'Get the book' CTA used at the end of every freebie."""
    shp(slide, Cm(2), top, Cm(17), Cm(6), CREAM, MSO_SHAPE.ROUNDED_RECTANGLE)
    # Gold accents
    shp(slide, Cm(2.5), top + Cm(0.2), Cm(1.5), Cm(0.1), GOLD, MSO_SHAPE.RECTANGLE)
    shp(slide, Cm(17), top + Cm(5.7), Cm(1.5), Cm(0.1), GOLD, MSO_SHAPE.RECTANGLE)

    txt(slide, Cm(2), top + Cm(0.5), Cm(17), Cm(1),
        title, Pt(22), NAVY, b=True, al=PP_ALIGN.CENTER, fn=FONT_HEADING)
    txt(slide, Cm(2), top + Cm(1.7), Cm(17), Cm(0.6),
        subtitle, Pt(11), SLATE, i=True, al=PP_ALIGN.CENTER)

    # Button
    shp(slide, Cm(6), top + Cm(2.7), Cm(9), Cm(1.3), NAVY, MSO_SHAPE.ROUNDED_RECTANGLE)
    txt(slide, Cm(6), top + Cm(2.9), Cm(9), Cm(0.9),
        button_text, Pt(13), GOLD, b=True, al=PP_ALIGN.CENTER, anc=MSO_ANCHOR.MIDDLE)

    # URL inside the box
    txt(slide, Cm(2), top + Cm(4.3), Cm(17), Cm(0.5),
        url, Pt(10), NAVY, b=True, al=PP_ALIGN.CENTER)


def cover_decoration(slide):
    """The standard navy cover decoration (circles + gold bars)."""
    shp(slide, Cm(15), Cm(-3), Cm(12), Cm(12), NAVY_LITE, MSO_SHAPE.OVAL)
    shp(slide, Cm(-4), Cm(20), Cm(8), Cm(8), NAVY_MID, MSO_SHAPE.OVAL)
    shp(slide, Cm(0), Cm(0), SW, Cm(0.5), GOLD, MSO_SHAPE.RECTANGLE)
    shp(slide, Cm(0), Cm(29.2), SW, Cm(0.5), GOLD, MSO_SHAPE.RECTANGLE)
