"""
Bundle Index — START HERE (PowerPoint Premium)
"Sleep, Baby. Please." by Six & Thriving (2026)
─────────────────────────────────────────────────────────────────────
Updated for the curated 7-file premium bundle.
2-page guide: what's in your bundle + where to start by situation.

Run: python bundle_index.py
"""

import os
from pptx import Presentation
from pptx.util import Pt, Cm
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from _brand import (
    NAVY, NAVY_MID, NAVY_LITE, GOLD, GOLD_LITE, CREAM, CREAM_MID, WHITE,
    TEAL, SLATE, SLATE_LITE, RUST, GREEN_DK, AMBER, AMBER_LT, BLUE_LT, GREEN_LT,
    PURPLE,
    FONT_HEADING, FONT_BODY, SW, SH,
    bg, shp, txt, mtxt, gbar, header, footer, dark_footer, callout,
    cta_box, cover_decoration,
)

OUTPUT = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                      "Bundle_Index_StartHere.pptx")


def p1_welcome(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg(slide, CREAM)
    header(slide, "START HERE")
    footer(slide, 1)

    txt(slide, Cm(2), Cm(2.2), Cm(15), Cm(0.4),
        "READ THIS FIRST", Pt(7), GOLD, b=True)
    gbar(slide, Cm(2), Cm(2.7), Cm(17))

    txt(slide, Cm(2), Cm(3.2), Cm(17), Cm(1.5),
        "Welcome to Your\nSleep Bundle",
        Pt(28), NAVY, b=True, fn=FONT_HEADING)
    txt(slide, Cm(2), Cm(6), Cm(17), Cm(0.6),
        "Everything you need to get your baby sleeping — in the right order.",
        Pt(12), GOLD, i=True)

    # Intro callout
    callout(slide, Cm(7),
            "DON'T FEEL OVERWHELMED",
            "You have 7 files. Each one has a clear purpose. This page tells you exactly what to open first — based on YOUR situation tonight.",
            accent=TEAL, bg_clr=BLUE_LT, h=Cm(2.8), text_color=NAVY)

    # The 7 files
    txt(slide, Cm(2), Cm(10.5), Cm(17), Cm(0.8),
        "YOUR BUNDLE INCLUDES",
        Pt(13), NAVY, b=True, fn=FONT_HEADING)
    txt(slide, Cm(12), Cm(10.5), Cm(7), Cm(0.6),
        "7 premium tools", Pt(10), GOLD, i=True, al=PP_ALIGN.RIGHT)
    gbar(slide, Cm(2), Cm(11.3), Cm(17))

    items = [
        ("01", "The Main eBook", "Sleep, Baby. Please. — 59 pages, 15 chapters"),
        ("02", "7-Night Sleep Tracker", "Fill in each morning. See the progress by Day 4."),
        ("03", "Bedtime Routine Builder", "Write it once. Stick it on the door. Use every night."),
        ("04", "Six Sleep Personalities Kit", "Find YOUR baby's type in 5 minutes."),
        ("05", "Reset Protocol Worksheet", "For parents who tried before. Start here."),
        ("06", "Mom Referral Card", "Share with a friend who needs this."),
        ("07", "This Guide (Start Here)", "You're reading it. You're in the right place."),
    ]
    y = Cm(12)
    for num, title, sub in items:
        card = shp(slide, Cm(2), y, Cm(17), Cm(1.8), WHITE, MSO_SHAPE.ROUNDED_RECTANGLE)
        card.line.color.rgb = CREAM_MID
        card.line.width = Pt(0.5)
        # Number badge
        shp(slide, Cm(2.4), y + Cm(0.35), Cm(1.1), Cm(1.1), NAVY, MSO_SHAPE.OVAL)
        txt(slide, Cm(2.4), y + Cm(0.4), Cm(1.1), Cm(1),
            num, Pt(9), GOLD, b=True, al=PP_ALIGN.CENTER, anc=MSO_ANCHOR.MIDDLE)
        # Title
        txt(slide, Cm(3.8), y + Cm(0.2), Cm(14), Cm(0.6),
            title, Pt(11), NAVY, b=True)
        # Subtitle
        txt(slide, Cm(3.8), y + Cm(0.9), Cm(14), Cm(0.6),
            sub, Pt(9.5), SLATE, i=True)
        y += Cm(2)

    # Bottom note
    callout(slide, Cm(26),
            "THAT'S IT. 7 FILES.",
            "No overwhelm. No filler. Each one earns its place. Open the ebook first — everything else supports it.",
            accent=GOLD, bg_clr=AMBER_LT, h=Cm(1.8), text_color=NAVY)


def p2_where_to_start(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg(slide, CREAM)
    header(slide, "Where to Start")
    footer(slide, 2)

    txt(slide, Cm(2), Cm(2.2), Cm(15), Cm(0.4),
        "PICK YOUR SITUATION", Pt(7), GOLD, b=True)
    gbar(slide, Cm(2), Cm(2.7), Cm(17))

    txt(slide, Cm(2), Cm(3.2), Cm(17), Cm(1.5),
        "Where to Start Tonight",
        Pt(24), NAVY, b=True, fn=FONT_HEADING)
    txt(slide, Cm(2), Cm(5), Cm(17), Cm(0.6),
        "Find the card that sounds like you. Do those steps.",
        Pt(11), GOLD, i=True)

    # 4 situation cards (2x2 grid)
    cards = [
        ("IT'S 2AM RIGHT NOW", RUST,
         "1. Read the Fast-Track page (page 5 of the ebook)\n2. Do the 6 steps tonight\n3. Read the rest tomorrow when you've slept"),
        ("I HAVE A PLAN — STARTING TONIGHT", TEAL,
         "1. Read Chapters 1–4 of the ebook\n2. Fill in the Bedtime Routine Builder\n3. Start the 7-Night Tracker tomorrow morning"),
        ("I'VE TRIED BEFORE — IT FAILED", AMBER,
         "1. Open the Reset Protocol Worksheet\n2. Diagnose your root cause (5 options)\n3. Read Chapter 7 of the ebook\n4. Commit to 5 nights"),
        ("I WANT TO KNOW MY BABY'S TYPE", PURPLE,
         "1. Open the Six Sleep Personalities Kit\n2. Find your baby's type (5 minutes)\n3. Read the matching strategy in Chapter 8"),
    ]

    positions = [
        (Cm(2), Cm(6.2)),      # top-left
        (Cm(10.5), Cm(6.2)),   # top-right
        (Cm(2), Cm(15.5)),     # bottom-left
        (Cm(10.5), Cm(15.5)),  # bottom-right
    ]
    card_w = Cm(8.2)
    card_h = Cm(8.5)

    for (title, color, steps), (cx, cy) in zip(cards, positions):
        # Card background
        card = shp(slide, cx, cy, card_w, card_h, WHITE, MSO_SHAPE.ROUNDED_RECTANGLE)
        card.line.color.rgb = color
        card.line.width = Pt(1.2)
        # Color header strip
        shp(slide, cx, cy, card_w, Cm(1.5), color, MSO_SHAPE.ROUNDED_RECTANGLE)
        shp(slide, cx, cy + Cm(1.3), card_w, Cm(0.3), color, MSO_SHAPE.RECTANGLE)
        # Title in header
        txt(slide, cx, cy + Cm(0.3), card_w, Cm(0.9),
            title, Pt(10), WHITE, b=True, al=PP_ALIGN.CENTER)
        # Steps
        txt(slide, cx + Cm(0.5), cy + Cm(2.2), card_w - Cm(1), Cm(6),
            steps, Pt(9.5), SLATE)

    # Bottom callout
    callout(slide, Cm(25),
            "ONE THING AT A TIME",
            "Don't try to use everything at once. Pick ONE card. Do those steps. Come back next week for the next layer.",
            accent=GOLD, bg_clr=AMBER_LT, h=Cm(2.5), text_color=NAVY)


def build():
    prs = Presentation()
    prs.slide_width = SW
    prs.slide_height = SH

    print()
    print("  ╔═══════════════════════════════════════════════╗")
    print("  ║  Bundle Index — START HERE (curated 7-file)   ║")
    print("  ╚═══════════════════════════════════════════════╝")
    print()

    pages = [
        ("Welcome + 7 files listed", p1_welcome),
        ("Where to start by situation", p2_where_to_start),
    ]
    for name, fn in pages:
        print(f"  ✓ {name}")
        fn(prs)

    prs.save(OUTPUT)
    size_kb = os.path.getsize(OUTPUT) / 1024
    print()
    print(f"  ✅ Saved: {OUTPUT}")
    print(f"  📄 Size: {size_kb:.1f} KB  •  2 pages")
    print()


if __name__ == "__main__":
    build()
