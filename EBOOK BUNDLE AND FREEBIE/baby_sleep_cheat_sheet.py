"""
The Baby Sleep Cheat Sheet — PREMIUM LEAD MAGNET (PowerPoint Edition)
"Sleep, Baby. Please." by Six & Thriving (2026)
─────────────────────────────────────────────────────────────────────
6-page emotional hook delivered as .pptx — same design system as the main ebook.
Editable in PowerPoint AND Canva. Export to PDF for distribution.

Page 1 — Branded cover with #1 BESTSELLER feel
Page 2 — "A letter to the mom reading this at 2am" (emotional hook)
Page 3 — The Awake Windows Master Table (the value)
Page 4 — 5 Reasons + Traffic Light Cues
Page 5 — Tonight's 6-Step Plan
Page 6 — The Night 3 Hook + book CTA

Run: python baby_sleep_cheat_sheet.py
"""

from pptx import Presentation
from pptx.util import Pt, Cm
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os

# ════════════════════════════════════════════════════════════════════
# BRAND PALETTE (identical to main ebook)
# ════════════════════════════════════════════════════════════════════
NAVY = RGBColor(0x0D, 0x1B, 0x3E)
NAVY_MID = RGBColor(0x16, 0x23, 0x47)
NAVY_LITE = RGBColor(0x1E, 0x2F, 0x55)
GOLD = RGBColor(0xC9, 0xA8, 0x4C)
GOLD_LITE = RGBColor(0xE8, 0xC9, 0x7A)
CREAM = RGBColor(0xFA, 0xF7, 0xF2)
CREAM_MID = RGBColor(0xF0, 0xEB, 0xE1)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TEAL = RGBColor(0x2D, 0x8A, 0x8A)
SLATE = RGBColor(0x4A, 0x55, 0x68)
SLATE_LITE = RGBColor(0x71, 0x80, 0x96)
RUST = RGBColor(0xC0, 0x39, 0x2B)
GREEN_DK = RGBColor(0x27, 0x67, 0x49)
AMBER = RGBColor(0xD6, 0x9E, 0x2E)
PURPLE = RGBColor(0x55, 0x3C, 0x9A)
AMBER_LT = RGBColor(0xFE, 0xFC, 0xBF)
BLUE_LT = RGBColor(0xEB, 0xF8, 0xFF)
RED_LT = RGBColor(0xFE, 0xD7, 0xD7)

# Fonts (matches main ebook)
FONT_HEADING = "Georgia"
FONT_BODY = "Arial"

# Page = A4 portrait (matches main ebook)
SW = Cm(21.0)
SH = Cm(29.7)

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT = os.path.join(SCRIPT_DIR, "Baby_Sleep_Cheat_Sheet.pptx")


# ════════════════════════════════════════════════════════════════════
# CORE HELPERS (same conventions as main ebook generator)
# ════════════════════════════════════════════════════════════════════

def bg(slide, color):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = color


def shp(slide, l, t, w, h, color, st=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(st, l, t, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    return s


def txt(slide, l, t, w, h, text, sz=Pt(10), clr=SLATE, b=False, i=False,
        al=PP_ALIGN.LEFT, fn=FONT_BODY, anc=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anc
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = sz
    p.font.color.rgb = clr
    p.font.bold = b
    p.font.italic = i
    p.font.name = fn
    p.alignment = al
    return tb


def mtxt(slide, l, t, w, h, paras):
    """Multi-paragraph textbox."""
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for idx, pd in enumerate(paras):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = pd.get("t", "")
        p.font.size = pd.get("sz", Pt(10))
        p.font.color.rgb = pd.get("c", SLATE)
        p.font.bold = pd.get("b", False)
        p.font.italic = pd.get("i", False)
        p.font.name = pd.get("fn", FONT_BODY)
        p.alignment = pd.get("al", PP_ALIGN.LEFT)
        p.space_after = pd.get("sa", Pt(6))
    return tb


def gbar(slide, l, t, w):
    return shp(slide, l, t, w, Cm(0.12), GOLD, MSO_SHAPE.RECTANGLE)


def header(slide, label_right):
    """Standard navy header band (matches main ebook)."""
    shp(slide, Cm(0), Cm(0), SW, Cm(1.3), NAVY, MSO_SHAPE.RECTANGLE)
    shp(slide, Cm(0), Cm(1.3), SW, Cm(0.12), GOLD, MSO_SHAPE.RECTANGLE)
    txt(slide, Cm(2), Cm(0.3), Cm(8), Cm(0.7),
        "Sleep, Baby. Please.", Pt(7.5), CREAM, fn=FONT_HEADING)
    txt(slide, Cm(11), Cm(0.3), Cm(8), Cm(0.7),
        label_right, Pt(7.5), GOLD_LITE, al=PP_ALIGN.RIGHT, fn=FONT_BODY)


def footer(slide, pg):
    """Standard cream footer with navy page-number circle."""
    shp(slide, Cm(0), Cm(28.1), SW, Cm(1.6), CREAM_MID, MSO_SHAPE.RECTANGLE)
    shp(slide, Cm(0), Cm(28.0), SW, Cm(0.1), GOLD, MSO_SHAPE.RECTANGLE)
    shp(slide, Cm(9.6), Cm(28.3), Cm(1.3), Cm(1.3), NAVY, MSO_SHAPE.OVAL)
    txt(slide, Cm(9.6), Cm(28.4), Cm(1.3), Cm(1.1), str(pg),
        Pt(9), WHITE, b=True, al=PP_ALIGN.CENTER, anc=MSO_ANCHOR.MIDDLE)
    txt(slide, Cm(2), Cm(28.4), Cm(7), Cm(0.7),
        "Six & Thriving  •  www.sixandthriving.com", Pt(7), SLATE_LITE)
    txt(slide, Cm(12.5), Cm(28.4), Cm(6.5), Cm(0.7),
        "© 2026 Six & Thriving", Pt(7), SLATE_LITE, al=PP_ALIGN.RIGHT)


def dark_footer(slide, pg):
    shp(slide, Cm(0), Cm(28.1), SW, Cm(1.6), RGBColor(0x08, 0x0F, 0x24), MSO_SHAPE.RECTANGLE)
    shp(slide, Cm(0), Cm(28.0), SW, Cm(0.1), GOLD, MSO_SHAPE.RECTANGLE)
    shp(slide, Cm(9.6), Cm(28.3), Cm(1.3), Cm(1.3), GOLD, MSO_SHAPE.OVAL)
    txt(slide, Cm(9.6), Cm(28.4), Cm(1.3), Cm(1.1), str(pg),
        Pt(9), NAVY, b=True, al=PP_ALIGN.CENTER, anc=MSO_ANCHOR.MIDDLE)
    txt(slide, Cm(2), Cm(28.4), Cm(7), Cm(0.7),
        "www.sixandthriving.com", Pt(7), GOLD_LITE)
    txt(slide, Cm(12.5), Cm(28.4), Cm(6.5), Cm(0.7),
        "© 2026 Six & Thriving", Pt(7), SLATE_LITE, al=PP_ALIGN.RIGHT)


# ════════════════════════════════════════════════════════════════════
# PAGE 1 — PREMIUM COVER
# ════════════════════════════════════════════════════════════════════

def p1_cover(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg(slide, NAVY)

    # Decorative circles
    shp(slide, Cm(15), Cm(-3), Cm(12), Cm(12), NAVY_LITE, MSO_SHAPE.OVAL)
    shp(slide, Cm(-4), Cm(20), Cm(8), Cm(8), NAVY_MID, MSO_SHAPE.OVAL)

    # Gold top bar
    shp(slide, Cm(0), Cm(0), SW, Cm(0.5), GOLD, MSO_SHAPE.RECTANGLE)
    # Gold bottom bar
    shp(slide, Cm(0), Cm(29.2), SW, Cm(0.5), GOLD, MSO_SHAPE.RECTANGLE)

    # FREE badge — gold circle, top-right
    shp(slide, Cm(15.5), Cm(2.5), Cm(3.5), Cm(3.5), GOLD, MSO_SHAPE.OVAL)
    txt(slide, Cm(15.5), Cm(3.2), Cm(3.5), Cm(0.9),
        "FREE", Pt(16), NAVY, b=True, al=PP_ALIGN.CENTER, fn=FONT_HEADING)
    txt(slide, Cm(15.5), Cm(4.2), Cm(3.5), Cm(0.7),
        "GUIDE", Pt(9), NAVY, b=True, al=PP_ALIGN.CENTER)

    # Series tag
    txt(slide, Cm(2), Cm(4.5), Cm(13), Cm(0.7),
        "FROM THE AUTHOR OF SLEEP, BABY. PLEASE.",
        Pt(9), GOLD, b=True, al=PP_ALIGN.LEFT)
    gbar(slide, Cm(2), Cm(5.3), Cm(7))

    # Main title
    txt(slide, Cm(2), Cm(8), Cm(17), Cm(2.5),
        "The Baby Sleep", Pt(48), WHITE, b=True, fn=FONT_HEADING)
    txt(slide, Cm(2), Cm(11), Cm(17), Cm(2.5),
        "Cheat Sheet", Pt(48), GOLD, b=True, i=True, fn=FONT_HEADING)

    # Diamond rule
    gbar(slide, Cm(2), Cm(14.3), Cm(7))
    shp(slide, Cm(5.3), Cm(14.15), Cm(0.4), Cm(0.4), GOLD, MSO_SHAPE.OVAL)

    # Subtitle
    mtxt(slide, Cm(2), Cm(15.2), Cm(17), Cm(2), [
        {"t": "Wake windows, nap counts, and bedtime targets",
         "sz": Pt(13), "c": CREAM_MID, "sa": Pt(4)},
        {"t": "for every age — birth to 24 months.",
         "sz": Pt(13), "c": CREAM_MID},
    ])

    # 3 promise bullets
    bullets = [
        "✦  The exact awake window for your baby's age",
        "✦  The 5 reasons your baby fights sleep",
        "✦  Tonight's 6-step plan you can start in 10 minutes",
    ]
    paras = [{"t": b, "sz": Pt(11), "c": GOLD_LITE, "sa": Pt(8)} for b in bullets]
    mtxt(slide, Cm(2), Cm(18), Cm(17), Cm(5), paras)

    # Author block
    gbar(slide, Cm(2), Cm(24), Cm(7))
    txt(slide, Cm(2), Cm(24.5), Cm(17), Cm(1),
        "SIX  &  THRIVING", Pt(16), WHITE, b=True, fn=FONT_HEADING)
    txt(slide, Cm(2), Cm(25.7), Cm(17), Cm(0.8),
        "Mother of Six — Still Here at 3 a.m.",
        Pt(11), GOLD_LITE, i=True)

    # Bottom URL
    txt(slide, Cm(0), Cm(28.5), SW, Cm(0.6),
        "www.sixandthriving.com",
        Pt(9), GOLD, al=PP_ALIGN.CENTER)


# ════════════════════════════════════════════════════════════════════
# PAGE 2 — EMOTIONAL HOOK LETTER
# ════════════════════════════════════════════════════════════════════

def p2_letter(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg(slide, CREAM)
    header(slide, "A Letter")
    footer(slide, 2)

    txt(slide, Cm(2), Cm(2.2), Cm(15), Cm(0.4),
        "BEFORE WE GET TO THE NUMBERS", Pt(7), GOLD, b=True)
    gbar(slide, Cm(2), Cm(2.7), Cm(17))

    txt(slide, Cm(2), Cm(3.2), Cm(17), Cm(2.5),
        "If you're reading this\nat 2 a.m. — this is for you.",
        Pt(26), NAVY, b=True, fn=FONT_HEADING)
    gbar(slide, Cm(2), Cm(7), Cm(3))

    body = [
        {"t": "I see you. Probably in a rocking chair. Maybe on the nursery floor. Holding a baby who should — by every logical measure — be exhausted. And yet here you both are.",
         "sz": Pt(11), "c": SLATE, "sa": Pt(10)},
        {"t": "I've been there. Six times.",
         "sz": Pt(11), "c": NAVY, "b": True, "sa": Pt(10)},
        {"t": "I want you to know something before you read another word:",
         "sz": Pt(11), "c": SLATE, "sa": Pt(10)},
    ]
    mtxt(slide, Cm(2), Cm(7.8), Cm(17), Cm(7), body)

    # Pull quote — navy box
    shp(slide, Cm(2), Cm(13.5), Cm(17), Cm(3.2), NAVY, MSO_SHAPE.ROUNDED_RECTANGLE)
    shp(slide, Cm(2), Cm(13.5), Cm(0.2), Cm(3.2), GOLD, MSO_SHAPE.RECTANGLE)
    mtxt(slide, Cm(2.6), Cm(13.7), Cm(16), Cm(3), [
        {"t": '"Your baby is not broken.', "sz": Pt(15), "c": GOLD_LITE,
         "b": True, "i": True, "fn": FONT_HEADING, "sa": Pt(2)},
        {"t": "You are not failing.", "sz": Pt(15), "c": GOLD_LITE,
         "b": True, "i": True, "fn": FONT_HEADING, "sa": Pt(4)},
        {"t": 'This is temporary, solvable, and you will come through it."',
         "sz": Pt(11), "c": CREAM, "i": True},
    ])

    closing = [
        {"t": "I'm a mom of six — including a set of twins. I've spent the last decade in the trenches of infant and toddler sleep. I'm not a sleep clinic. I don't have a PhD. I lived it. Every scenario in the full guide.",
         "sz": Pt(11), "c": SLATE, "sa": Pt(10)},
        {"t": "And every single one of my babies learned to sleep. Every single one.",
         "sz": Pt(11), "c": NAVY, "b": True, "sa": Pt(14)},
    ]
    mtxt(slide, Cm(2), Cm(17.5), Cm(17), Cm(5), closing)

    # Sign-off
    mtxt(slide, Cm(2), Cm(22.5), Cm(17), Cm(2), [
        {"t": "With love and solidarity,", "sz": Pt(10), "c": SLATE, "i": True, "sa": Pt(4)},
        {"t": "— Six & Thriving", "sz": Pt(11), "c": GOLD, "b": True, "fn": FONT_HEADING},
    ])

    # Tease for next page (amber)
    shp(slide, Cm(2), Cm(25), Cm(17), Cm(2), AMBER_LT, MSO_SHAPE.ROUNDED_RECTANGLE)
    shp(slide, Cm(2), Cm(25), Cm(0.2), Cm(2), AMBER, MSO_SHAPE.RECTANGLE)
    txt(slide, Cm(2.6), Cm(25.2), Cm(16), Cm(0.4),
        "TURN THE PAGE →", Pt(8), AMBER, b=True)
    txt(slide, Cm(2.6), Cm(25.8), Cm(16), Cm(1),
        "The awake windows table. This is what you came for.",
        Pt(10), NAVY)


# ════════════════════════════════════════════════════════════════════
# PAGE 3 — AWAKE WINDOWS MASTER TABLE
# ════════════════════════════════════════════════════════════════════

def p3_table(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg(slide, CREAM)
    header(slide, "Awake Windows")
    footer(slide, 3)

    txt(slide, Cm(2), Cm(2.2), Cm(15), Cm(0.4),
        "REFERENCE TABLE", Pt(7), GOLD, b=True)
    gbar(slide, Cm(2), Cm(2.7), Cm(17))

    txt(slide, Cm(2), Cm(3.1), Cm(17), Cm(1.5),
        "Awake Windows by Age", Pt(24), NAVY, b=True, fn=FONT_HEADING)
    txt(slide, Cm(2), Cm(5.1), Cm(17), Cm(0.6),
        "Birth to 24 months — at a glance.",
        Pt(11), GOLD, i=True)

    # Table
    rows = [
        ("AGE", "AWAKE WINDOW", "NAPS/DAY", "BEDTIME"),
        ("0–6 weeks", "45–60 min", "4–5", "Variable"),
        ("6–12 weeks", "60–90 min", "4", "8–9 pm"),
        ("3–4 months", "90 min – 2 hrs", "3–4", "7:30–8:30 pm"),
        ("4–6 months", "2 – 2.5 hrs", "3", "7–8 pm"),
        ("6–8 months", "2.5 – 3 hrs", "2–3", "7–8 pm"),
        ("8–12 months", "3 – 3.5 hrs", "2", "6:30–7:30 pm"),
        ("12–18 months", "4 – 5 hrs", "1–2", "7–7:30 pm"),
        ("18–24 months", "5 – 6 hrs", "1", "7–7:30 pm"),
    ]
    col_widths = [Cm(3.5), Cm(4.5), Cm(3), Cm(6)]
    y = Cm(6.5)
    row_h = Cm(1.2)

    for i, row in enumerate(rows):
        if i == 0:
            shp(slide, Cm(2), y, Cm(17), row_h, NAVY, MSO_SHAPE.RECTANGLE)
            x = Cm(2)
            for j, cell in enumerate(row):
                txt(slide, x, y + Cm(0.3), col_widths[j], Cm(0.7),
                    cell, Pt(10), WHITE, b=True, al=PP_ALIGN.CENTER)
                x += col_widths[j]
            y += row_h
        else:
            row_color = WHITE if i % 2 == 1 else CREAM_MID
            row_bg = shp(slide, Cm(2), y, Cm(17), row_h, row_color, MSO_SHAPE.RECTANGLE)
            row_bg.line.color.rgb = RGBColor(0xD4, 0xB8, 0x96)
            row_bg.line.width = Pt(0.3)
            x = Cm(2)
            for j, cell in enumerate(row):
                color = NAVY if j == 0 else SLATE
                bold = j == 0
                txt(slide, x, y + Cm(0.35), col_widths[j], Cm(0.6),
                    cell, Pt(10), color, b=bold, al=PP_ALIGN.CENTER)
                x += col_widths[j]
            y += row_h

    # The #1 mistake callout
    shp(slide, Cm(2), Cm(18.5), Cm(17), Cm(3.2), AMBER_LT, MSO_SHAPE.ROUNDED_RECTANGLE)
    shp(slide, Cm(2), Cm(18.5), Cm(0.2), Cm(3.2), AMBER, MSO_SHAPE.RECTANGLE)
    txt(slide, Cm(2.6), Cm(18.7), Cm(16), Cm(0.5),
        "THE #1 TIMING MISTAKE PARENTS MAKE", Pt(8), AMBER, b=True)
    txt(slide, Cm(2.6), Cm(19.4), Cm(16), Cm(2.5),
        "Putting baby down at a fixed clock time ('we always do 7pm') regardless of when their last nap ended. Always count FORWARD from wake time, not BACKWARD from a target.",
        Pt(10), NAVY)

    # Tease for next page
    shp(slide, Cm(2), Cm(22.5), Cm(17), Cm(2), BLUE_LT, MSO_SHAPE.ROUNDED_RECTANGLE)
    shp(slide, Cm(2), Cm(22.5), Cm(0.2), Cm(2), TEAL, MSO_SHAPE.RECTANGLE)
    txt(slide, Cm(2.6), Cm(22.7), Cm(16), Cm(0.4),
        "NEXT PAGE →", Pt(8), TEAL, b=True)
    txt(slide, Cm(2.6), Cm(23.3), Cm(16), Cm(1),
        "The 5 reasons your baby is fighting sleep — and how to read their cues.",
        Pt(10), NAVY)


# ════════════════════════════════════════════════════════════════════
# PAGE 4 — 5 REASONS + TRAFFIC LIGHT CUES
# ════════════════════════════════════════════════════════════════════

def p4_reasons(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg(slide, CREAM)
    header(slide, "Why Baby Isn't Sleeping")
    footer(slide, 4)

    txt(slide, Cm(2), Cm(2.2), Cm(15), Cm(0.4),
        "DIAGNOSE FIRST, FIX SECOND", Pt(7), GOLD, b=True)
    gbar(slide, Cm(2), Cm(2.7), Cm(17))

    txt(slide, Cm(2), Cm(3.1), Cm(17), Cm(1.2),
        "5 Reasons Babies Fight Sleep", Pt(22), NAVY, b=True, fn=FONT_HEADING)
    txt(slide, Cm(2), Cm(4.7), Cm(17), Cm(0.6),
        "Most struggles come from a surprisingly short list.",
        Pt(11), GOLD, i=True)

    # 5 reason cards
    reasons = [
        ("Sleep\nAssociations", NAVY),
        ("Over-\ntiredness", GREEN_DK),
        ("Under-\ntiredness", AMBER),
        ("Environment", TEAL),
        ("Developmental\nLeaps", PURPLE),
    ]
    card_w = Cm(3.2)
    gap = Cm(0.25)
    x = Cm(2)
    for label, color in reasons:
        shp(slide, x, Cm(5.8), card_w, Cm(2.7), color, MSO_SHAPE.ROUNDED_RECTANGLE)
        txt(slide, x, Cm(6.5), card_w, Cm(2),
            label, Pt(9.5), WHITE, b=True, al=PP_ALIGN.CENTER, fn=FONT_HEADING)
        x += card_w + gap

    # Sleep Associations deep dive
    txt(slide, Cm(2), Cm(9.3), Cm(17), Cm(0.8),
        "The #1 Hidden Driver: Sleep Associations",
        Pt(14), NAVY, b=True, fn=FONT_HEADING)
    txt(slide, Cm(2), Cm(10.3), Cm(17), Cm(2.5),
        "Whatever your baby falls asleep WITH, they will need at every wake-up. Always feed to sleep? They'll need to feed at 1am, 3am, 5am. Always rocked? They'll need rocking at every cycle. This is the #1 reason babies who 'used to sleep great' suddenly start waking — usually around 4 months. It is biology, not manipulation.",
        Pt(10), SLATE)

    # Traffic Light section
    txt(slide, Cm(2), Cm(13.3), Cm(17), Cm(0.8),
        "Tiredness Cues — The Traffic Light",
        Pt(14), NAVY, b=True, fn=FONT_HEADING)
    txt(slide, Cm(2), Cm(14.3), Cm(17), Cm(0.6),
        "Catch your baby at GREEN. By RED, cortisol has kicked in.",
        Pt(10), GOLD, i=True)

    # Three traffic cards
    stages = [
        ("GREEN LIGHT", GREEN_DK,
         ["Slowing down", "Quiet gaze", "Single yawn"], "ACT NOW"),
        ("YELLOW LIGHT", AMBER,
         ["Frequent yawning", "Eye rubbing", "Ear pulling"], "BEGIN ROUTINE"),
        ("RED LIGHT", RUST,
         ["Crying", "Arching back", "Wired & hyper"], "DAMAGE CONTROL"),
    ]
    tx = Cm(2)
    tw = Cm(5.5)
    tgap = Cm(0.25)
    for label, color, cues, cta in stages:
        shp(slide, tx, Cm(15.2), tw, Cm(7), color, MSO_SHAPE.ROUNDED_RECTANGLE)
        txt(slide, tx, Cm(15.5), tw, Cm(0.8),
            label, Pt(11), WHITE, b=True, al=PP_ALIGN.CENTER, fn=FONT_HEADING)
        # Cues
        cue_text = "\n".join(["• " + c for c in cues])
        txt(slide, tx + Cm(0.3), Cm(16.8), tw - Cm(0.6), Cm(3.5),
            cue_text, Pt(9.5), WHITE, al=PP_ALIGN.CENTER)
        # CTA pill
        shp(slide, tx + Cm(0.4), Cm(20.7), tw - Cm(0.8), Cm(1.2),
            NAVY_LITE, MSO_SHAPE.ROUNDED_RECTANGLE)
        txt(slide, tx + Cm(0.4), Cm(20.9), tw - Cm(0.8), Cm(0.8),
            cta, Pt(9), WHITE, b=True, al=PP_ALIGN.CENTER, anc=MSO_ANCHOR.MIDDLE)
        tx += tw + tgap

    # Bottom hook
    shp(slide, Cm(2), Cm(23), Cm(17), Cm(3), BLUE_LT, MSO_SHAPE.ROUNDED_RECTANGLE)
    shp(slide, Cm(2), Cm(23), Cm(0.2), Cm(3), TEAL, MSO_SHAPE.RECTANGLE)
    txt(slide, Cm(2.6), Cm(23.2), Cm(16), Cm(0.4),
        "STAGE 1 IS YOUR WINDOW", Pt(8), TEAL, b=True)
    txt(slide, Cm(2.6), Cm(23.8), Cm(16), Cm(2.2),
        "Most parents only catch RED — by then you're in damage control. The full book teaches you to catch GREEN every time. Turn the page for tonight's plan.",
        Pt(10), NAVY)


# ════════════════════════════════════════════════════════════════════
# PAGE 5 — TONIGHT'S 6-STEP PLAN
# ════════════════════════════════════════════════════════════════════

def p5_plan(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg(slide, CREAM)
    header(slide, "Tonight's Plan")
    footer(slide, 5)

    txt(slide, Cm(2), Cm(2.2), Cm(15), Cm(0.4),
        "DO THESE TONIGHT", Pt(7), GOLD, b=True)
    gbar(slide, Cm(2), Cm(2.7), Cm(17))

    txt(slide, Cm(2), Cm(3.1), Cm(17), Cm(1.2),
        "Tonight's 6-Step Plan", Pt(22), NAVY, b=True, fn=FONT_HEADING)
    txt(slide, Cm(2), Cm(4.7), Cm(17), Cm(0.6),
        "If you read nothing else, do these six things tonight.",
        Pt(11), GOLD, i=True)

    steps = [
        ("1", "Check the room",
         "Pitch black? White noise running? 68–72°F / 20–22°C?"),
        ("2", "Find the awake window",
         "Use the table on page 3. Set a timer."),
        ("3", "Write the bedtime routine — IN ORDER",
         "Bath → lotion → pyjamas → feed → song → crib awake."),
        ("4", "Decide your response plan BEFORE the night starts",
         "Your 2 a.m. brain won't make good decisions. Write it down."),
        ("5", "Pause before you go in",
         "Wait 2–3 minutes when you hear a sound. Many babies resettle."),
        ("6", "Crib AWAKE",
         "Drowsy is fine. Asleep is not. This single skill changes everything."),
    ]

    y = Cm(6)
    for num, title, body in steps:
        # White card
        card = shp(slide, Cm(2), y, Cm(17), Cm(2.6), WHITE, MSO_SHAPE.ROUNDED_RECTANGLE)
        card.line.color.rgb = CREAM_MID
        card.line.width = Pt(0.5)
        # Number circle
        shp(slide, Cm(2.4), y + Cm(0.6), Cm(1.4), Cm(1.4), NAVY, MSO_SHAPE.OVAL)
        txt(slide, Cm(2.4), y + Cm(0.7), Cm(1.4), Cm(1.2),
            num, Pt(13), GOLD, b=True, al=PP_ALIGN.CENTER, anc=MSO_ANCHOR.MIDDLE,
            fn=FONT_HEADING)
        # Title + body
        txt(slide, Cm(4.3), y + Cm(0.4), Cm(14), Cm(0.7),
            title, Pt(11), NAVY, b=True)
        txt(slide, Cm(4.3), y + Cm(1.3), Cm(14), Cm(1.2),
            body, Pt(9.5), SLATE)
        y += Cm(2.9)

    # Done is better than perfect
    shp(slide, Cm(2), Cm(24.5), Cm(17), Cm(2.5), AMBER_LT, MSO_SHAPE.ROUNDED_RECTANGLE)
    shp(slide, Cm(2), Cm(24.5), Cm(0.2), Cm(2.5), AMBER, MSO_SHAPE.RECTANGLE)
    txt(slide, Cm(2.6), Cm(24.7), Cm(16), Cm(0.4),
        "DONE IS BETTER THAN PERFECT", Pt(8), AMBER, b=True)
    txt(slide, Cm(2.6), Cm(25.3), Cm(16), Cm(2),
        "Pick ONE thing from this list and do it tonight. Progress starts with a single consistent step — not all six executed flawlessly.",
        Pt(10), NAVY)


# ════════════════════════════════════════════════════════════════════
# PAGE 6 — THE NIGHT 3 HOOK + CTA
# ════════════════════════════════════════════════════════════════════

def p6_cta(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg(slide, NAVY)

    # Decorative
    shp(slide, Cm(15), Cm(-3), Cm(12), Cm(12), NAVY_LITE, MSO_SHAPE.OVAL)
    shp(slide, Cm(-4), Cm(20), Cm(8), Cm(8), NAVY_MID, MSO_SHAPE.OVAL)

    # Gold top + bottom
    shp(slide, Cm(0), Cm(0), SW, Cm(0.4), GOLD, MSO_SHAPE.RECTANGLE)
    shp(slide, Cm(0), Cm(29.3), SW, Cm(0.4), GOLD, MSO_SHAPE.RECTANGLE)

    # Section label
    txt(slide, Cm(2), Cm(3), Cm(17), Cm(0.6),
        "ONE LAST THING — THE MOST IMPORTANT THING",
        Pt(9), GOLD, b=True, al=PP_ALIGN.CENTER)

    # The hook
    txt(slide, Cm(2), Cm(5.5), Cm(17), Cm(2.5),
        "Nobody talks", Pt(40), WHITE, b=True, al=PP_ALIGN.CENTER, fn=FONT_HEADING)
    txt(slide, Cm(2), Cm(8), Cm(17), Cm(2.5),
        "about Night 3.", Pt(40), GOLD, b=True, i=True, al=PP_ALIGN.CENTER, fn=FONT_HEADING)

    # Gold rule
    gbar(slide, Cm(8.5), Cm(11), Cm(4))

    # Body
    body = [
        {"t": "Night 3 is the hardest night of sleep training. Not because the method is failing — because it is WORKING. When a previously reinforced behavior stops being rewarded, it escalates dramatically before it disappears.",
         "sz": Pt(11.5), "c": CREAM_MID, "sa": Pt(8), "al": PP_ALIGN.CENTER},
        {"t": "It's called the extinction burst. Every other guide either skips it or buries it. The full book names it, maps it night-by-night, and walks you through it.",
         "sz": Pt(11.5), "c": CREAM_MID, "sa": Pt(14), "al": PP_ALIGN.CENTER},
        {"t": "It's the most important thing nobody told you.",
         "sz": Pt(13), "c": GOLD_LITE, "b": True, "i": True, "fn": FONT_HEADING, "al": PP_ALIGN.CENTER},
    ]
    mtxt(slide, Cm(2), Cm(11.8), Cm(17), Cm(7), body)

    # === CTA BOX ===
    shp(slide, Cm(2), Cm(19.5), Cm(17), Cm(6), CREAM, MSO_SHAPE.ROUNDED_RECTANGLE)
    # Gold corner accents
    shp(slide, Cm(2.5), Cm(19.7), Cm(1.5), Cm(0.1), GOLD, MSO_SHAPE.RECTANGLE)
    shp(slide, Cm(17), Cm(25.3), Cm(1.5), Cm(0.1), GOLD, MSO_SHAPE.RECTANGLE)

    txt(slide, Cm(2), Cm(20), Cm(17), Cm(1),
        "Want the full plan?", Pt(22), NAVY, b=True, al=PP_ALIGN.CENTER, fn=FONT_HEADING)
    txt(slide, Cm(2), Cm(21.2), Cm(17), Cm(0.6),
        "Sleep, Baby. Please. — the 59-page evidence-based guide",
        Pt(11), SLATE, i=True, al=PP_ALIGN.CENTER)

    # 3 bullets
    bullets = [
        "✦  The full 7-night plan",
        "✦  Every method explained",
        "✦  The Six Sleep Personalities",
    ]
    bw = Cm(5.5)
    bx = Cm(2.2)
    for b in bullets:
        txt(slide, bx, Cm(22.2), bw, Cm(0.6),
            b, Pt(10), NAVY, al=PP_ALIGN.CENTER)
        bx += bw + Cm(0.15)

    # The button
    shp(slide, Cm(6), Cm(23.3), Cm(9), Cm(1.3), NAVY, MSO_SHAPE.ROUNDED_RECTANGLE)
    txt(slide, Cm(6), Cm(23.5), Cm(9), Cm(0.9),
        "GET THE BOOK  →  $19", Pt(13), GOLD, b=True, al=PP_ALIGN.CENTER, anc=MSO_ANCHOR.MIDDLE)

    # URL — INSIDE the cream box (this was the bug — was bleeding out)
    txt(slide, Cm(2), Cm(24.85), Cm(17), Cm(0.5),
        "www.sixandthriving.com",
        Pt(10), NAVY, b=True, al=PP_ALIGN.CENTER)

    # Final sign-off — ABOVE the dark footer (was overlapping URL)
    txt(slide, Cm(2), Cm(26.5), Cm(17), Cm(0.6),
        "Sleep is coming. For both of you.",
        Pt(11), GOLD_LITE, i=True, al=PP_ALIGN.CENTER, fn=FONT_HEADING)
    txt(slide, Cm(2), Cm(27.2), Cm(17), Cm(0.5),
        "— Six & Thriving",
        Pt(10), GOLD, b=True, al=PP_ALIGN.CENTER, fn=FONT_HEADING)

    dark_footer(slide, 6)


# ════════════════════════════════════════════════════════════════════
# BUILD
# ════════════════════════════════════════════════════════════════════

def build():
    prs = Presentation()
    prs.slide_width = SW
    prs.slide_height = SH

    print()
    print("  ╔═════════════════════════════════════════════════════════════╗")
    print("  ║  THE BABY SLEEP CHEAT SHEET — PREMIUM (.pptx) v3            ║")
    print("  ║  Six & Thriving | 2026                                      ║")
    print("  ╚═════════════════════════════════════════════════════════════╝")
    print()

    pages = [
        ("Cover", p1_cover),
        ("Letter to mom at 2am", p2_letter),
        ("Awake Windows table", p3_table),
        ("5 Reasons + Traffic Light", p4_reasons),
        ("Tonight's 6-Step Plan", p5_plan),
        ("Night 3 hook + CTA", p6_cta),
    ]

    for name, fn in pages:
        print(f"  ✓ {name}")
        fn(prs)

    prs.save(OUTPUT)
    size_kb = os.path.getsize(OUTPUT) / 1024

    print()
    print("  ═════════════════════════════════════════════════════════════")
    print(f"  ✅ Saved: {OUTPUT}")
    print(f"  📄 Size: {size_kb:.1f} KB  •  6 pages  •  A4 portrait")
    print()
    print("  📋 TO USE:")
    print("     • Open in PowerPoint to edit")
    print("     • OR upload to Canva (Create → Import file)")
    print("     • Export as PDF when ready to distribute")
    print()


if __name__ == "__main__":
    build()
